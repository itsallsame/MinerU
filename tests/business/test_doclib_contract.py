"""Live Doclib contract probe with an isolated home and no remote model access."""

from __future__ import annotations

import os
import io
import json
import shutil
import sqlite3
import subprocess
import sys
import tempfile
import time
from collections.abc import Iterator
from contextlib import closing
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from reportlab.pdfgen import canvas
from fastapi.testclient import TestClient

from mineru.business.api import create_app
from mineru.business.documents import DoclibGateway, ImmutableUploadStore
from mineru.business.services import BusinessDiscovery, DocumentWorkflow, EvidenceReader, EvidenceWriter, FieldExtraction
from mineru.business.store import BusinessStore
from mineru.doclib import DoclibClient, ParseReleaseRequest, ParseRequest, ScanRequest
from mineru.doclib.endpoint import read_endpoint_file
from mineru.doclib.services.parse_svc import parse_batch_json_path
from mineru.doclib.types import ForgetPathRequest
from mineru.errors import MineruError, ServerNotRunningError


@pytest.fixture
def live_doclib(tmp_path: Path) -> Iterator[tuple[DoclibClient, Path, Path]]:
    # macOS AF_UNIX paths are short; pytest's nested tmp_path is too long.
    with tempfile.TemporaryDirectory(prefix="mu-doclib-", dir="/private/tmp") as short_home:
        home = Path(short_home)
        socket_path = home / "doclib.sock"
        env = os.environ.copy()
        env.update(
            {
                "MINERU_HOME": str(home),
                "MINERU_MODEL_SOURCE": "local",
                "MINERU_MODEL_BASE_DIR": str(tmp_path / "models"),
                "MINERU_DOCLIB_UDS_ENABLED": "true",
                "MINERU_DOCLIB_TCP_ENABLED": "false",
                "HF_HUB_OFFLINE": "1",
                "TRANSFORMERS_OFFLINE": "1",
            }
        )
        log_path = tmp_path / "doclib-process.log"
        with log_path.open("w", encoding="utf-8") as log:
            process = subprocess.Popen(
                [sys.executable, "-m", "mineru.doclib.app"],
                env=env,
                stdout=log,
                stderr=subprocess.STDOUT,
            )
            client = DoclibClient(socket_path=socket_path, timeout=10)
            try:
                deadline = time.monotonic() + 30
                while time.monotonic() < deadline:
                    if process.poll() is not None:
                        raise RuntimeError(f"Doclib exited early: {log_path.read_text()}")
                    try:
                        client.get_server_status()
                        break
                    except (ServerNotRunningError, MineruError):
                        time.sleep(0.1)
                else:
                    raise RuntimeError(f"Doclib did not start: {log_path.read_text()}")
                yield client, tmp_path, home
            finally:
                client.close()
                process.terminate()
                try:
                    process.wait(timeout=10)
                except subprocess.TimeoutExpired:
                    process.kill()
                    process.wait(timeout=10)


def _wait_for_scan(client: DoclibClient, scan_id: int) -> None:
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline:
        scan = client.get_scan(scan_id)
        if scan.status == "done":
            return
        if scan.status == "failed":
            raise AssertionError(f"Doclib scan failed: {scan.error_code} {scan.error_msg}")
        time.sleep(0.1)
    raise AssertionError("Doclib scan timed out")


def _wait_for_parse(client: DoclibClient, parse_ids: list[int]) -> None:
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        states = [client.get_parse(parse_id) for parse_id in parse_ids]
        if all(state.status == "done" for state in states):
            return
        failures = [state for state in states if state.status == "failed"]
        if failures:
            raise AssertionError(f"Doclib parse failed: {[(state.error_code, state.error_msg) for state in failures]}")
        time.sleep(0.1)
    raise AssertionError("Doclib parse timed out")


def test_public_doclib_parse_registers_business_consumer(
    live_doclib: tuple[DoclibClient, Path, Path]
) -> None:
    client, root, home = live_doclib
    source = root / "consumer.html"
    source.write_text("<h1>Consumer registration</h1>", encoding="utf-8")

    result = client.ensure_parse(ParseRequest(path=str(source), tier="flash", consumer_key="business:contract"))
    assert len(result.wait_parse_ids) == 1
    with closing(sqlite3.connect(home / "doclib.db")) as database:
        claims = database.execute(
            "SELECT consumer_key, protected FROM parse_consumers WHERE parse_id=?",
            (result.wait_parse_ids[0],),
        ).fetchall()
    assert ("business:contract", 0) in claims
    assert set(claims) <= {("business:contract", 0), ("system:ingest", 1)}


def test_public_doclib_release_is_idempotent_and_rejects_late_requeue(
    live_doclib: tuple[DoclibClient, Path, Path]
) -> None:
    client, root, _home = live_doclib
    source = root / "release.html"
    source.write_text("<h1>Public release</h1>", encoding="utf-8")
    submitted = client.ensure_parse(
        ParseRequest(path=str(source), tier="flash", consumer_key="business:public-release")
    )
    released = client.release_parse_consumer(ParseReleaseRequest(consumer_key="business:public-release"))
    assert [item.parse_id for item in released.results] == submitted.wait_parse_ids
    assert released.results[0].disposition in {"skipped", "shared", "running", "finished", "retained"}
    assert client.release_parse_consumer(ParseReleaseRequest(consumer_key="business:public-release")) == released
    with pytest.raises(MineruError) as error:
        client.ensure_parse(ParseRequest(path=str(source), tier="flash", consumer_key="business:public-release"))
    assert error.value.code == "consumer_released"


def test_real_doclib_business_cancel_route_persists_release_facts(
    live_doclib: tuple[DoclibClient, Path, Path]
) -> None:
    doclib, root, _home = live_doclib
    upload_root = root / "business-uploads"
    upload_root.mkdir()
    uploads = ImmutableUploadStore(upload_root, max_bytes=1024)
    store = BusinessStore(root / "cancel-business.sqlite3")
    store.initialize()
    workflow = DocumentWorkflow(
        uploads=uploads, store=store, gateway=DoclibGateway(doclib, shared_root=upload_root),
        doclib=doclib, producer_version="4.0.6",
    )
    api = TestClient(create_app(
        workflow=workflow, store=store, uploads=uploads,
        evidence_reader=EvidenceReader(store=store, doclib=doclib),
        evidence_writer=EvidenceWriter(store=store, doclib=doclib),
    ))
    submitted = api.post(
        "/api/business/documents", files={"file": ("cancel.html", b"<h1>Live cancellation</h1>", "text/html")},
    )
    assert submitted.status_code == 202, submitted.text
    task_id = submitted.json()["task"]["id"]
    cancelled = api.post(f"/api/business/tasks/{task_id}/cancel")
    assert cancelled.status_code == 200, cancelled.text
    assert cancelled.json()["status"] == "cancelled"
    assert cancelled.json()["cancel_effect"] in {"queued_skipped", "may_continue"}
    assert api.post(f"/api/business/tasks/{task_id}/cancel").json() == cancelled.json()
    assert store.list_revisions(submitted.json()["document"]["id"]) == ()
    task = store.get_task(task_id)
    assert task is not None and task.cancel_results_json is not None


def test_native_html_doclib_round_trip(live_doclib: tuple[DoclibClient, Path, Path]) -> None:
    client, root, _home = live_doclib
    source = root / "report.html"
    source.write_text("<html><body><h1>Project Lantern</h1><p>Revenue increased by 12 percent.</p></body></html>")

    submitted = DoclibGateway(client, shared_root=root).submit(source)
    _wait_for_parse(client, list(submitted.parse_ids))
    doc = client.get_doc_by_path(str(source))
    assert len(doc.sha256) == 64
    assert submitted.sha256 == doc.sha256
    assert submitted.tier == "flash"
    content = client.get_doc_content(doc.sha256, tier="flash")
    assert "Project Lantern" in content.content
    assert content.tier == "flash"
    assert content.content_ranges
    locator = content.content_ranges[0].start
    recalled = client.read_content(locator)
    assert recalled.sha256 == doc.sha256
    assert "Project Lantern" in recalled.content

    forced = client.ensure_parse(ParseRequest(path=str(source), tier="flash", force=True, remote=False))
    _wait_for_parse(client, forced.created_parse_ids)
    assert forced.created_parse_ids
    assert not set(forced.created_parse_ids) & set(submitted.parse_ids)
    # Public locators encode source/tier/page/block, not parse_id. They are
    # navigation hints; business evidence must keep its own revision snapshot.
    assert "/parse:" not in locator

    search = client.search("Lantern", tier="flash")
    assert any(result.sha256 == doc.sha256 for result in search.results)


def test_business_discovery_uses_real_local_doclib_without_exposing_paths(
    live_doclib: tuple[DoclibClient, Path, Path]
) -> None:
    client, root, _home = live_doclib
    source = root / "business-lantern.html"
    source.write_text("<html><body><h1>DistinctiveLanternProject</h1></body></html>", encoding="utf-8")
    submitted = DoclibGateway(client, shared_root=root).submit(source)
    _wait_for_parse(client, list(submitted.parse_ids))
    business_dir = root / "business-discovery"
    business_dir.mkdir()
    store = BusinessStore(business_dir / "business.sqlite3")
    store.initialize()
    upload = ImmutableUploadStore(root, max_bytes=1024).store(
        io.BytesIO(source.read_bytes()), filename="business-lantern.html"
    )
    document = store.create_document(upload, original_name="business-lantern.html")
    revision = store.add_completed_revision(
        document.id, parse=client.get_parse(submitted.parse_ids[0]), producer_version="4.0.6"
    )
    discovery = BusinessDiscovery(store=store, doclib=client)
    matches = discovery.search("DistinctiveLanternProject")
    assert any(hit.document.id == document.id and hit.revision_id == revision.id for hit in matches.items)
    assert all("/private/" not in hit.snippet for hit in matches.items)
    content = client.get_doc_content(submitted.sha256, tier="flash")
    locator = content.content_ranges[0].start
    historical = discovery.read(revision.id, locator)
    assert historical.document_id == document.id
    assert "DistinctiveLanternProject" in historical.content
    outline = discovery.outline(revision.id)
    assert outline.scanned_pages == 1 and outline.next_page is None
    assert any(item.title == "DistinctiveLanternProject" and item.level == 1 for item in outline.items)
    assert all(item.locator.startswith(f"doc:{revision.short_id}/tier:flash/page:") for item in outline.items)
    native_blocks = client.read_parse_structure(submitted.parse_ids[0], 1)
    assert native_blocks.sha256 == submitted.sha256 and native_blocks.blocks
    assert any(block.preview == "DistinctiveLanternProject" for block in native_blocks.blocks)
    structure = discovery.structure(revision.id, 1)
    assert structure.document_id == document.id
    assert any(block.locator.startswith(f"doc:{revision.short_id}/tier:flash/page:1") for block in structure.blocks)


def test_business_field_candidate_uses_historical_doclib_page(
    live_doclib: tuple[DoclibClient, Path, Path]
) -> None:
    client, root, home = live_doclib
    source = root / "notice.html"
    source.write_text(
        "<html><body><p>标题：年度通知</p><p>发文单位：办公室</p></body></html>", encoding="utf-8"
    )
    submitted = DoclibGateway(client, shared_root=root).submit(source)
    _wait_for_parse(client, list(submitted.parse_ids))
    business_dir = root / "business-field-test"
    business_dir.mkdir()
    business = BusinessStore(business_dir / "business.sqlite3")
    business.initialize()
    stored = ImmutableUploadStore(root, max_bytes=1024).store(
        io.BytesIO(source.read_bytes()), filename="notice.html"
    )
    document = business.create_document(stored, original_name="notice.html", template_code="official_document")
    parse = client.get_parse(submitted.parse_ids[0])
    revision = business.add_completed_revision(document.id, parse=parse, producer_version="4.0.6")
    time.sleep(0.02)
    current = client.ensure_parse(ParseRequest(path=str(source), tier="flash", force=True, remote=False))
    _wait_for_parse(client, current.created_parse_ids)
    current_parse = client.get_parse(current.created_parse_ids[0])
    current_json = Path(
        parse_batch_json_path(
            str(home / "doclib"), submitted.sha256, "flash", current_parse.page_range, current_parse.done_at
        )
    )
    payload = current_json.read_text()
    assert "年度通知" in payload
    current_json.write_text(payload.replace("年度通知", "当前通知"))
    writer = EvidenceWriter(store=business, doclib=client)
    extraction = FieldExtraction(store=business, doclib=client, evidence_writer=writer)
    extraction.enqueue(revision.id)
    run = extraction.process_next()
    assert run is not None
    assert run.status == "done"
    candidates = business.list_field_candidates(run.id)
    assert {(item.field_code, item.value) for item in candidates} == {
        ("title", "年度通知"), ("issuer", "办公室")
    }
    assert all(business.get_evidence(item.evidence_id).revision_id == revision.id for item in candidates)
    locator = f"doc:{submitted.short_id}/tier:flash/page:1"
    assert "当前通知" in client.read_content(locator).content
    assert "年度通知" in client.read_parse_content(parse.id, locator).content


def test_parse_id_content_read_does_not_substitute_a_newer_batch(live_doclib: tuple[DoclibClient, Path, Path]) -> None:
    client, root, home = live_doclib
    source = root / "revision.html"
    source.write_text("<h1>Historical Lantern</h1>")
    first = DoclibGateway(client, shared_root=root).submit(source)
    _wait_for_parse(client, list(first.parse_ids))
    first_id = first.parse_ids[0]
    locator = client.get_doc_content(first.sha256, tier="flash").content_ranges[0].start
    assert "Historical Lantern" in client.read_parse_content(first_id, locator).content

    time.sleep(0.02)
    second = client.ensure_parse(ParseRequest(path=str(source), tier="flash", force=True, remote=False))
    _wait_for_parse(client, second.created_parse_ids)
    second_id = second.created_parse_ids[0]
    first_info = client.get_parse(first_id)
    second_info = client.get_parse(second_id)
    first_path = Path(
        parse_batch_json_path(str(home / "doclib"), first.sha256, "flash", first_info.page_range, first_info.done_at)
    )
    second_path = Path(
        parse_batch_json_path(str(home / "doclib"), first.sha256, "flash", second_info.page_range, second_info.done_at)
    )
    assert first_path != second_path
    assert first_path.is_file() and second_path.is_file()
    payload = second_path.read_text()
    assert "Historical Lantern" in payload
    second_path.write_text(payload.replace("Historical Lantern", "Current Lantern"))

    assert "Historical Lantern" in client.read_parse_content(first_id, locator).content
    assert "Current Lantern" in client.read_parse_content(second_id, locator).content
    assert "Current Lantern" in client.read_content(locator).content
    assert any(block.preview == "Historical Lantern" for block in client.read_parse_structure(first_id, 1).blocks)
    assert any(block.preview == "Current Lantern" for block in client.read_parse_structure(second_id, 1).blocks)
    historical_hits = client.search_parse_blocks(first_id, 1, "Historical Lantern")
    assert len(historical_hits.matches) == 1
    assert historical_hits.matches[0].block_no == 1
    assert "Historical Lantern" in historical_hits.matches[0].snippet
    assert client.search_parse_blocks(second_id, 1, "Historical Lantern").matches == []

    business_dir = root / "business-revisions"
    business_dir.mkdir()
    business = BusinessStore(business_dir / "business.sqlite3")
    business.initialize()
    upload = ImmutableUploadStore(root, max_bytes=1024).store(io.BytesIO(source.read_bytes()), filename="revision.html")
    document = business.create_document(upload, original_name="revision.html")
    revision = business.add_completed_revision(document.id, parse=first_info, producer_version="4.0.6")
    second_revision = business.add_completed_revision(document.id, parse=second_info, producer_version="4.0.7")
    diff = BusinessDiscovery(store=business, doclib=client).diff_revisions(revision.id, second_revision.id)
    assert [item.status for item in diff.items] == ["changed"]
    assert diff.items[0].left_locator == diff.items[0].right_locator == locator
    block_page = BusinessDiscovery(store=business, doclib=client).search_blocks(revision.id, "Historical Lantern")
    assert len(block_page.items) == 1
    assert block_page.items[0].locator == historical_hits.matches[0].locator
    frozen_block = EvidenceWriter(store=business, doclib=client).capture(
        revision.id, locator=block_page.items[0].locator,
    )
    assert "Historical Lantern" in frozen_block.snippet
    assert "Current Lantern" not in frozen_block.snippet
    evidence = EvidenceWriter(store=business, doclib=client).capture(revision.id, locator=locator)
    assert "Historical Lantern" in evidence.snippet
    assert "Current Lantern" not in evidence.snippet

    with pytest.raises(MineruError, match="cached"):
        client.read_parse_content(999999, locator)
    with pytest.raises(MineruError, match="does not belong"):
        client.read_parse_content(first_id, f"doc:{first_info.short_id}/tier:basic/page:1/block:1")


def test_native_structure_preserves_nested_model_blocks_without_child_locator_claim(
    live_doclib: tuple[DoclibClient, Path, Path],
) -> None:
    client, root, home = live_doclib
    source = root / "nested.html"
    source.write_text("<h1>Native tree</h1>", encoding="utf-8")
    submitted = DoclibGateway(client, shared_root=root).submit(source)
    _wait_for_parse(client, list(submitted.parse_ids))
    parse = client.get_parse(submitted.parse_ids[0])
    middle_path = Path(parse_batch_json_path(
        str(home / "doclib"), submitted.sha256, "flash", parse.page_range, parse.done_at,
    ))
    middle = json.loads(middle_path.read_text(encoding="utf-8"))
    middle["pages"][0]["blocks"] = [{
        "type": "list", "index": 0,
        "content": [{"type": "text", "content": [{"type": "text", "content": "Nested item"}]}],
    }]
    middle_path.write_text(json.dumps(middle), encoding="utf-8")
    tree = client.read_parse_structure(parse.id, 1)
    assert len(tree.blocks) == 1
    parent = tree.blocks[0]
    assert parent.type == "list" and parent.path == [0]
    assert len(parent.children) == 1
    child = parent.children[0]
    assert child.type == "text" and child.path == [0, 0] and child.preview == "Nested item"
    assert child.locator == parent.locator and child.block_no == parent.block_no
    matches = client.search_parse_blocks(parse.id, 1, "Nested item")
    assert len(matches.matches) == 1
    assert matches.matches[0].locator == parent.locator


def test_published_upload_can_be_submitted_to_doclib(live_doclib: tuple[DoclibClient, Path, Path]) -> None:
    client, root, _home = live_doclib
    stored = ImmutableUploadStore(root, max_bytes=1024).store(
        io.BytesIO(b"<h1>Project Lantern published source</h1>"), filename="report.html"
    )
    submitted = DoclibGateway(client, shared_root=root).submit(stored.path)
    _wait_for_parse(client, list(submitted.parse_ids))
    assert submitted.sha256 == stored.sha256
    assert "Project Lantern published" in client.get_doc_content(submitted.sha256, tier="flash").content


def test_business_evidence_references_real_doclib_parse(live_doclib: tuple[DoclibClient, Path, Path]) -> None:
    client, root, _home = live_doclib
    stored = ImmutableUploadStore(root, max_bytes=1024).store(
        io.BytesIO(b"<h1>Project Lantern real evidence</h1>"), filename="report.html"
    )
    database_dir = root / "business"
    database_dir.mkdir()
    business = BusinessStore(database_dir / "business.sqlite3")
    business.initialize()
    document = business.create_document(stored, original_name="report.html")
    submitted = DoclibGateway(client, shared_root=root).submit(stored.path)
    _wait_for_parse(client, list(submitted.parse_ids))
    parse = client.get_parse(submitted.parse_ids[0])
    revision = business.add_completed_revision(document.id, parse=parse, producer_version="4.0.6")
    content = client.get_doc_content(submitted.sha256, tier="flash")
    locator = content.content_ranges[0].start
    snippet = client.read_content(locator).content
    evidence = business.capture_evidence(revision.id, locator=locator, snippet=snippet)

    assert evidence.document_id == document.id
    assert evidence.snippet == snippet
    assert BusinessStore(database_dir / "business.sqlite3").get_evidence(evidence.id) == evidence
    assert EvidenceReader(store=business, doclib=client).inspect(evidence.id).navigation_status == "current_match"


def test_doclib_content_identity_is_hash_based(live_doclib: tuple[DoclibClient, Path, Path]) -> None:
    client, root, _home = live_doclib
    first = root / "a" / "same.html"
    duplicate = root / "b" / "same.html"
    different = root / "c" / "same.html"
    for path in (first, duplicate, different):
        path.parent.mkdir()
    first.write_text("<html><body>Identical document</body></html>")
    duplicate.write_bytes(first.read_bytes())
    different.write_text("<html><body>Different document</body></html>")

    for path in (first, duplicate, different):
        scan = client.create_scan(ScanRequest(path=str(path), kind="manual", source="sdk"))
        _wait_for_scan(client, scan.id)

    first_doc = client.get_doc_by_path(str(first))
    duplicate_doc = client.get_doc_by_path(str(duplicate))
    different_doc = client.get_doc_by_path(str(different))
    assert first_doc.sha256 == duplicate_doc.sha256
    assert first_doc.short_id == duplicate_doc.short_id
    assert first_doc.sha256 != different_doc.sha256


def test_same_content_reuses_trackable_parse_id(live_doclib: tuple[DoclibClient, Path, Path]) -> None:
    client, root, _home = live_doclib
    first = root / "first.html"
    second = root / "second.html"
    first.write_text("<h1>Project Lantern duplicate parse</h1>")
    second.write_bytes(first.read_bytes())
    gateway = DoclibGateway(client, shared_root=root)
    submitted = gateway.submit(first)
    _wait_for_parse(client, list(submitted.parse_ids))
    reused = gateway.submit(second)
    assert reused.sha256 == submitted.sha256
    assert reused.parse_ids
    assert all(client.get_parse(parse_id).status == "done" for parse_id in reused.parse_ids)


def test_text_pdf_flash_parse_and_locator(live_doclib: tuple[DoclibClient, Path, Path]) -> None:
    client, root, _home = live_doclib
    source = root / "report.pdf"
    pdf = canvas.Canvas(str(source))
    for index in range(20):
        pdf.drawString(72, 720 - index * 24, f"Project Lantern revenue 12 percent in quarter {index + 1}.")
    pdf.save()

    submitted = DoclibGateway(client, shared_root=root).submit(source, tier="flash")
    _wait_for_parse(client, list(submitted.parse_ids))
    content = client.get_doc_content(submitted.sha256, tier="flash")
    assert "Project Lantern" in content.content
    assert content.content_ranges
    assert client.read_content(content.content_ranges[0].start).sha256 == submitted.sha256


def test_repo_paper_pdf_round_trip_through_real_business_api_and_doclib(
    live_doclib: tuple[DoclibClient, Path, Path]
) -> None:
    doclib, root, _home = live_doclib
    sample = Path(__file__).resolve().parents[2] / "demo" / "pdfs" / "demo1.pdf"
    assert sample.is_file()
    upload_root = root / "business-uploads"
    upload_root.mkdir()
    uploads = ImmutableUploadStore(upload_root, max_bytes=5 * 1024 * 1024)
    store = BusinessStore(root / "paper-business.sqlite3")
    store.initialize()
    workflow = DocumentWorkflow(
        uploads=uploads, store=store, gateway=DoclibGateway(doclib, shared_root=upload_root),
        doclib=doclib, producer_version="4.0.6",
    )
    client = TestClient(create_app(
        workflow=workflow, store=store, uploads=uploads,
        evidence_reader=EvidenceReader(store=store, doclib=doclib),
        evidence_writer=EvidenceWriter(store=store, doclib=doclib),
        discovery=BusinessDiscovery(store=store, doclib=doclib),
    ))
    with sample.open("rb") as source:
        uploaded = client.post("/api/business/documents", files={"file": (sample.name, source, "application/pdf")},
                               data={"tier": "flash"})
    assert uploaded.status_code == 202, uploaded.text
    payload = uploaded.json()
    document = payload["document"]
    task_id = payload["task"]["id"]
    assert payload["task"]["status"] == "submitted"
    assert client.get(f"/api/business/documents/{document['id']}/source").content == sample.read_bytes()

    deadline = time.monotonic() + 45
    while time.monotonic() < deadline:
        task = client.get(f"/api/business/tasks/{task_id}").json()
        if task["status"] in ("done", "failed"):
            break
        time.sleep(0.2)
    else:
        raise AssertionError("Repository PDF business parse did not reach a terminal state")
    if task["status"] != "done":
        saved_task = store.get_task(task_id)
        assert saved_task is not None
        parses = [doclib.get_parse(parse_id) for parse_id in saved_task.parse_ids]
        details = [(parse.id, parse.status, parse.page_range, parse.tier) for parse in parses]
        raise AssertionError(f"Business task failed: {task['error_code']}; parse batches: {details}")
    revisions = client.get(f"/api/business/documents/{document['id']}/revisions").json()
    assert len(revisions) == 1 and revisions[0]["tier"] == "flash"
    assert revisions[0]["page_range"] == "1-13"
    locator = f"doc:{revisions[0]['short_id']}/tier:flash/page:1"
    historical = client.get(f"/api/business/revisions/{revisions[0]['id']}/content", params={"locator": locator})
    assert historical.status_code == 200, historical.text
    assert "afforestation" in historical.json()["content"].lower()
    assert historical.json()["state"] == "historical_parse_unconfirmed"
    last_page = client.get(
        f"/api/business/revisions/{revisions[0]['id']}/content",
        params={"locator": f"doc:{revisions[0]['short_id']}/tier:flash/page:13"},
    )
    assert last_page.status_code == 200, last_page.text
    page_matches = client.get(
        f"/api/business/revisions/{revisions[0]['id']}/search", params={"query": "afforestation"},
    )
    assert page_matches.status_code == 200, page_matches.text
    assert page_matches.json()["items"]
    assert all(item["locator"].startswith(f"doc:{revisions[0]['short_id']}/tier:flash/page:")
               for item in page_matches.json()["items"])
    assert page_matches.json()["next_page"] is None
    outline = client.get(f"/api/business/revisions/{revisions[0]['id']}/outline")
    assert outline.status_code == 200, outline.text
    assert outline.json()["scanned_pages"] == 13 and outline.json()["next_page"] is None
    assert all(item["locator"].startswith(f"doc:{revisions[0]['short_id']}/tier:flash/page:")
               for item in outline.json()["items"])
    last_structure = client.get(
        f"/api/business/revisions/{revisions[0]['id']}/structure", params={"page_no": 13},
    )
    assert last_structure.status_code == 200, last_structure.text
    assert last_structure.json()["page_no"] == 13
    assert all(item["locator"].startswith(f"doc:{revisions[0]['short_id']}/tier:flash/page:13")
               for item in last_structure.json()["blocks"])
    assert "parse_id" not in last_structure.text

    captured = client.post(f"/api/business/revisions/{revisions[0]['id']}/evidence", json={"locator": locator})
    assert captured.status_code == 201, captured.text
    inspected = client.get(f"/api/business/evidence/{captured.json()['id']}")
    assert inspected.status_code == 200
    assert inspected.json()["snippet"] == captured.json()["snippet"]
    assert inspected.json()["navigation_status"] == "current_match"
    searched = client.get("/api/business/search", params={"query": "afforestation"})
    assert searched.status_code == 200, searched.text
    assert any(item["document"]["id"] == document["id"] for item in searched.json()["items"])

    # Identical bytes create a separate business document but reuse the local
    # Doclib parse without leaking its batch IDs through the public API.
    with sample.open("rb") as source:
        duplicate = client.post("/api/business/documents", files={"file": (sample.name, source, "application/pdf")},
                                data={"tier": "flash"})
    assert duplicate.status_code == 202, duplicate.text
    duplicate_doc = duplicate.json()["document"]
    assert duplicate_doc["id"] != document["id"]
    duplicate_task = client.get(f"/api/business/tasks/{duplicate.json()['task']['id']}").json()
    if duplicate_task["status"] != "done":
        saved_duplicate = store.get_task(duplicate.json()["task"]["id"])
        assert saved_duplicate is not None
        batch_states = [
            (parse_id, doclib.get_parse(parse_id).status, doclib.get_parse(parse_id).page_range)
            for parse_id in saved_duplicate.parse_ids
        ]
        raise AssertionError(f"Duplicate task failed: {duplicate_task['error_code']}; batches={batch_states}")
    duplicate_revisions = client.get(f"/api/business/documents/{duplicate_doc['id']}/revisions").json()
    assert len(duplicate_revisions) == 1 and duplicate_revisions[0]["page_range"] == "1-13"


@pytest.mark.parametrize("extension", ["docx", "pptx", "xlsx"])
def test_native_office_flash_parse(live_doclib: tuple[DoclibClient, Path, Path], extension: str) -> None:
    client, root, _home = live_doclib
    source = root / f"report.{extension}"
    if extension == "docx":
        document = Document()
        document.add_paragraph("Project Lantern office evidence")
        document.save(source)
    elif extension == "pptx":
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        textbox = slide.shapes.add_textbox(0, 0, 3_000_000, 500_000)
        textbox.text = "Project Lantern office evidence"
        presentation.save(source)
    else:
        workbook = Workbook()
        workbook.active["A1"] = "Project Lantern office evidence"
        workbook.save(source)

    submitted = DoclibGateway(client, shared_root=root).submit(source)
    _wait_for_parse(client, list(submitted.parse_ids))
    content = client.get_doc_content(submitted.sha256, tier="flash")
    assert submitted.tier == "flash"
    assert "Project Lantern" in content.content


def test_image_ingest_reports_missing_local_model(live_doclib: tuple[DoclibClient, Path, Path]) -> None:
    client, root, _home = live_doclib
    source = root / "scan.png"
    Image.new("RGB", (200, 100), "white").save(source)

    submitted = DoclibGateway(client, shared_root=root).submit(source, tier="flash")
    doc = client.get_doc_by_path(str(source))
    assert doc.sha256 == submitted.sha256

    deadline = time.monotonic() + 15
    while time.monotonic() < deadline:
        state = client.get_parse(submitted.parse_ids[0])
        if state.status in ("done", "failed"):
            break
        time.sleep(0.1)
    else:
        raise AssertionError("Image parse did not reach a terminal state")
    assert state.status == "failed"
    assert state.error_code == "parse_failed"
    assert "not ready" in (state.error_msg or "")


def test_forget_path_can_be_reingested_while_source_exists(live_doclib: tuple[DoclibClient, Path, Path]) -> None:
    client, root, _home = live_doclib
    source = root / "report.html"
    source.write_text("<h1>Project Lantern immutable source</h1>")
    submitted = DoclibGateway(client, shared_root=root).submit(source)
    _wait_for_parse(client, list(submitted.parse_ids))

    preview = client.forget_path(ForgetPathRequest(path=str(source), dry_run=True))
    assert preview.forgotten_files == 1
    assert client.get_doc_by_path(str(source)).sha256 == submitted.sha256

    result = client.forget_path(ForgetPathRequest(path=str(source), dry_run=False))
    assert result.forgotten_files == 1
    assert not any(file.path == str(source) for file in client.list_files().files)
    # get_doc_by_path auto-ingests a still-existing file; forgetting the path is
    # not a durable business deletion policy.
    assert client.get_doc_by_path(str(source)).sha256 == submitted.sha256
    assert "Project Lantern" in client.get_doc_content(submitted.sha256, tier="flash").content


def test_doclib_restart_preserves_content_and_locator(live_doclib: tuple[DoclibClient, Path, Path]) -> None:
    client, root, home = live_doclib
    source = root / "report.html"
    source.write_text("<h1>Project Lantern restart evidence</h1>")
    submitted = DoclibGateway(client, shared_root=root).submit(source)
    _wait_for_parse(client, list(submitted.parse_ids))
    locator = client.get_doc_content(submitted.sha256, tier="flash").content_ranges[0].start
    assert client.shutdown_server().accepted
    deadline = time.monotonic() + 10
    while (home / "doclib.sock").exists() and time.monotonic() < deadline:
        time.sleep(0.1)

    env = os.environ.copy()
    env.update(
        {
            "MINERU_HOME": str(home),
            "MINERU_MODEL_SOURCE": "local",
            "MINERU_MODEL_BASE_DIR": str(root / "models"),
            "MINERU_DOCLIB_UDS_ENABLED": "true",
            "MINERU_DOCLIB_TCP_ENABLED": "false",
            "HF_HUB_OFFLINE": "1",
            "TRANSFORMERS_OFFLINE": "1",
        }
    )
    with (root / "restart.log").open("w", encoding="utf-8") as log:
        restarted = subprocess.Popen([sys.executable, "-m", "mineru.doclib.app"], env=env, stdout=log, stderr=subprocess.STDOUT)
        fresh_client = DoclibClient(socket_path=home / "doclib.sock", timeout=10)
        try:
            deadline = time.monotonic() + 20
            while time.monotonic() < deadline:
                if restarted.poll() is not None:
                    raise AssertionError(f"Restarted Doclib exited: {(root / 'restart.log').read_text()}")
                try:
                    fresh_client.get_server_status()
                    break
                except (ServerNotRunningError, MineruError):
                    time.sleep(0.1)
            else:
                raise AssertionError("Restarted Doclib did not become ready")
            assert fresh_client.get_doc(submitted.sha256).sha256 == submitted.sha256
            assert "Project Lantern" in fresh_client.read_content(locator).content
        finally:
            fresh_client.close()
            restarted.terminate()
            try:
                restarted.wait(timeout=10)
            except subprocess.TimeoutExpired:
                restarted.kill()
                restarted.wait(timeout=10)


def test_explicit_tcp_client_reaches_isolated_doclib(tmp_path: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="mu-doclib-tcp-", dir="/private/tmp") as short_home:
        home = Path(short_home)
        env = os.environ.copy()
        env.update(
            {
                "MINERU_HOME": str(home),
                "MINERU_MODEL_SOURCE": "local",
                "MINERU_MODEL_BASE_DIR": str(tmp_path / "models"),
                "MINERU_DOCLIB_UDS_ENABLED": "false",
                "MINERU_DOCLIB_TCP_ENABLED": "true",
                "MINERU_DOCLIB_TCP_HOST": "127.0.0.1",
                "MINERU_DOCLIB_TCP_PORT": "0",
                "HF_HUB_OFFLINE": "1",
                "TRANSFORMERS_OFFLINE": "1",
            }
        )
        log_path = tmp_path / "doclib-tcp.log"
        with log_path.open("w", encoding="utf-8") as log:
            process = subprocess.Popen(
                [sys.executable, "-m", "mineru.doclib.app"], env=env, stdout=log, stderr=subprocess.STDOUT
            )
            client: DoclibClient | None = None
            try:
                deadline = time.monotonic() + 30
                while time.monotonic() < deadline:
                    if process.poll() is not None:
                        raise AssertionError(f"TCP Doclib exited: {log_path.read_text()}")
                    endpoint = read_endpoint_file(home / "doclib.endpoint.json")
                    urls = (
                        [transport.base_url for transport in endpoint.transports if transport.type == "tcp"] if endpoint else []
                    )
                    if urls and urls[0]:
                        client = DoclibClient(base_url=urls[0], timeout=10)
                        try:
                            client.get_server_status()
                            break
                        except (ServerNotRunningError, MineruError):
                            client.close()
                            client = None
                    time.sleep(0.1)
                else:
                    raise AssertionError(f"TCP Doclib did not start: {log_path.read_text()}")

                source = tmp_path / "tcp-report.html"
                source.write_text("<h1>Project Lantern TCP evidence</h1>")
                submitted = DoclibGateway(client, shared_root=tmp_path).submit(source)
                _wait_for_parse(client, list(submitted.parse_ids))
                assert "Project Lantern TCP" in client.get_doc_content(submitted.sha256, tier="flash").content
            finally:
                if client is not None:
                    client.close()
                process.terminate()
                try:
                    process.wait(timeout=10)
                except subprocess.TimeoutExpired:
                    process.kill()
                    process.wait(timeout=10)


def test_second_doclib_cannot_own_same_home(live_doclib: tuple[DoclibClient, Path, Path]) -> None:
    client, root, home = live_doclib
    env = os.environ.copy()
    env.update(
        {
            "MINERU_HOME": str(home),
            "MINERU_MODEL_SOURCE": "local",
            "MINERU_MODEL_BASE_DIR": str(root / "models"),
            "MINERU_DOCLIB_UDS_ENABLED": "true",
            "MINERU_DOCLIB_TCP_ENABLED": "false",
            "HF_HUB_OFFLINE": "1",
            "TRANSFORMERS_OFFLINE": "1",
        }
    )
    competing = subprocess.run(
        [sys.executable, "-m", "mineru.doclib.app"],
        env=env,
        capture_output=True,
        text=True,
        timeout=15,
        check=False,
    )
    assert competing.returncode != 0
    assert "currently owned by another doclib server process" in competing.stdout + competing.stderr
    assert client.get_server_status() is not None


def test_offline_backup_restores_parsed_content(live_doclib: tuple[DoclibClient, Path, Path]) -> None:
    client, root, home = live_doclib
    source = root / "backup-report.html"
    source.write_text("<h1>Project Lantern backup evidence</h1>")
    submitted = DoclibGateway(client, shared_root=root).submit(source)
    _wait_for_parse(client, list(submitted.parse_ids))
    locator = client.get_doc_content(submitted.sha256, tier="flash").content_ranges[0].start
    assert client.shutdown_server().accepted
    deadline = time.monotonic() + 10
    while (home / "doclib.sock").exists() and time.monotonic() < deadline:
        time.sleep(0.1)
    assert not (home / "doclib.sock").exists()

    with tempfile.TemporaryDirectory(prefix="mu-doclib-restore-", dir="/private/tmp") as restore_root:
        restored_home = Path(restore_root) / "home"
        shutil.copytree(home, restored_home)
        env = os.environ.copy()
        env.update(
            {
                "MINERU_HOME": str(restored_home),
                "MINERU_MODEL_SOURCE": "local",
                "MINERU_MODEL_BASE_DIR": str(root / "models"),
                "MINERU_DOCLIB_UDS_ENABLED": "true",
                "MINERU_DOCLIB_TCP_ENABLED": "false",
                "HF_HUB_OFFLINE": "1",
                "TRANSFORMERS_OFFLINE": "1",
            }
        )
        log_path = root / "restore.log"
        with log_path.open("w", encoding="utf-8") as log:
            restored = subprocess.Popen(
                [sys.executable, "-m", "mineru.doclib.app"], env=env, stdout=log, stderr=subprocess.STDOUT
            )
            restored_client = DoclibClient(socket_path=restored_home / "doclib.sock", timeout=10)
            try:
                deadline = time.monotonic() + 20
                while time.monotonic() < deadline:
                    if restored.poll() is not None:
                        raise AssertionError(f"Restored Doclib exited: {log_path.read_text()}")
                    try:
                        restored_client.get_server_status()
                        break
                    except (ServerNotRunningError, MineruError):
                        time.sleep(0.1)
                else:
                    raise AssertionError(f"Restored Doclib did not start: {log_path.read_text()}")
                assert restored_client.get_doc(submitted.sha256).sha256 == submitted.sha256
                assert "Project Lantern backup" in restored_client.read_content(locator).content
            finally:
                restored_client.close()
                restored.terminate()
                try:
                    restored.wait(timeout=10)
                except subprocess.TimeoutExpired:
                    restored.kill()
                    restored.wait(timeout=10)
