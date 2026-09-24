"""Opt-in real Web/API/Doclib integration using the repository's public PDF sample."""

from __future__ import annotations

import hashlib
import json
import os
import socket
import subprocess
import sys
import tempfile
import time
from pathlib import Path
from urllib.error import URLError
from urllib.request import urlopen

import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from mineru.doclib import DoclibClient
from mineru.doclib.endpoint import read_endpoint_file


ROOT = Path(__file__).resolve().parents[2]
SAMPLE = ROOT / "demo" / "pdfs" / "demo1.pdf"
WEB_ROOT = ROOT / "business-web" / "dist"
SKILL_SCRIPT = ROOT / "skills" / "business-documents" / "scripts" / "business_documents.py"


def _wait_for_doclib(process: subprocess.Popen[bytes], home: Path, log_path: Path) -> str:
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        if process.poll() is not None:
            raise AssertionError(f"Doclib exited early: {log_path.read_text()}")
        endpoint = read_endpoint_file(home / "doclib.endpoint.json")
        urls = [item.base_url for item in endpoint.transports if item.type == "tcp"] if endpoint else []
        if urls and urls[0]:
            client = DoclibClient(base_url=urls[0], timeout=5)
            try:
                client.get_server_status()
                return urls[0]
            except Exception:
                pass
            finally:
                client.close()
        time.sleep(0.1)
    raise AssertionError(f"Doclib did not become ready: {log_path.read_text()}")


def _wait_for_api(process: subprocess.Popen[bytes], origin: str, log_path: Path) -> None:
    deadline = time.monotonic() + 30
    while time.monotonic() < deadline:
        if process.poll() is not None:
            raise AssertionError(f"Business API exited early: {log_path.read_text()}")
        try:
            with urlopen(f"{origin}/api/business/capabilities", timeout=2) as response:
                if response.status == 200:
                    return
        except (OSError, URLError):
            pass
        time.sleep(0.1)
    raise AssertionError(f"Business API did not become ready: {log_path.read_text()}")


def _stop(process: subprocess.Popen[bytes]) -> None:
    process.terminate()
    try:
        process.wait(timeout=10)
    except subprocess.TimeoutExpired:
        process.kill()
        process.wait(timeout=10)


def _skill(origin: str, *args: str) -> object:
    result = subprocess.run(
        [sys.executable, str(SKILL_SCRIPT), "--base-url", origin, *args],
        capture_output=True,
        text=True,
        timeout=30,
    )
    assert result.returncode == 0, f"Skill {args[0]} failed: {result.stderr.strip()}"
    return json.loads(result.stdout)


def _verify_live_skill(origin: str) -> None:
    uploaded = _skill(origin, "upload", str(SAMPLE), "--tier", "flash", "--template", "paper")
    assert isinstance(uploaded, dict)
    document = uploaded["document"]
    assert document["sha256"] == hashlib.sha256(SAMPLE.read_bytes()).hexdigest()
    assert document["template_code"] == "paper"
    task_id = uploaded["task"]["id"]
    deadline = time.monotonic() + 90
    while time.monotonic() < deadline:
        task = _skill(origin, "task", task_id)
        assert isinstance(task, dict)
        if task["status"] in ("done", "failed"):
            break
        time.sleep(0.5)
    else:
        raise AssertionError("Skill-uploaded PDF did not reach a terminal task status")
    assert task["status"] == "done", task
    overview = _skill(origin, "overview", document["id"])
    assert isinstance(overview, dict)
    revision = overview["revision"]
    assert revision["tier"] == "flash"
    assert overview["draft"] is None
    assert overview["confirmed_for_latest_run"] is None
    locator = f"doc:{revision['short_id']}/tier:{revision['tier']}/page:1"
    read = _skill(origin, "read", revision["id"], locator)
    assert isinstance(read, dict)
    assert read["state"] == "historical_parse_unconfirmed"
    assert read["content"].strip()
    started = _skill(origin, "extract", revision["id"])
    assert isinstance(started, dict)
    run_id = started["id"]
    deadline = time.monotonic() + 90
    while time.monotonic() < deadline:
        extraction = _skill(origin, "extraction", run_id)
        assert isinstance(extraction, dict)
        if extraction["run"]["status"] in ("done", "failed"):
            break
        time.sleep(0.5)
    else:
        raise AssertionError("Skill extraction did not reach a terminal status")
    assert extraction["state"] == "machine_unconfirmed"
    assert extraction["run"]["status"] == "done", extraction
    results = _skill(origin, "results", run_id)
    assert isinstance(results, dict)
    assert results == {"state": "confirmed", "items": []}


@pytest.mark.skipif(os.getenv("MINERU_RUN_LIVE_BROWSER") != "1", reason="Opt in to local Chromium and socket integration")
def test_live_web_native_formats_through_business_api_and_doclib() -> None:
    assert SAMPLE.is_file()
    assert hashlib.sha256(SAMPLE.read_bytes()).hexdigest() == "f3b3be345bf2df8979f2491ca9466e078e4fd1d6a216611faa8566e4c44d474b"
    assert (WEB_ROOT / "asset-manifest.json").is_file(), "Run cd business-web && pnpm build first"
    temporary_root = "/private/tmp" if sys.platform == "darwin" else None
    with tempfile.TemporaryDirectory(prefix="mu-live-web-", dir=temporary_root) as temporary:
        home = Path(temporary)
        uploads = home / "uploads"
        uploads.mkdir()
        office_files = [home / f"smoke.{extension}" for extension in ("docx", "pptx", "xlsx")]
        document = Document()
        document.add_paragraph("题目：MinerUOfficeDocxMarker")
        document.save(office_files[0])
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, 3_000_000, 500_000).text = "MinerUOfficePptxMarker"
        presentation.save(office_files[1])
        workbook = Workbook()
        workbook.active["A1"] = "MinerUOfficeXlsxMarker"
        workbook.save(office_files[2])
        doclib_log = home / "doclib.log"
        api_log = home / "business-api.log"
        doclib_env = os.environ.copy()
        doclib_env.update(
            {
                "MINERU_HOME": str(home / "doclib-home"),
                "MINERU_MODEL_SOURCE": "local",
                "MINERU_MODEL_BASE_DIR": str(home / "models"),
                "MINERU_DOCLIB_UDS_ENABLED": "false",
                "MINERU_DOCLIB_TCP_ENABLED": "true",
                "MINERU_DOCLIB_TCP_HOST": "127.0.0.1",
                "MINERU_DOCLIB_TCP_PORT": "0",
                "MINERU_MODEL_VLM_SERVER_URL": "",
                "MINERU_LLM_AIDED_FEATURES_TITLE_LEVELING": "false",
                "MINERU_LLM_AIDED_FEATURES_CROSS_PAGE_TABLE_CELL_MERGE": "false",
                "HF_HUB_OFFLINE": "1",
                "TRANSFORMERS_OFFLINE": "1",
                "HF_DATASETS_OFFLINE": "1",
            }
        )
        (home / "doclib-home").mkdir()
        with doclib_log.open("wb") as doclib_stream:
            doclib = subprocess.Popen(
                [sys.executable, "-m", "mineru.doclib.app"],
                env=doclib_env,
                stdout=doclib_stream,
                stderr=subprocess.STDOUT,
            )
        try:
            doclib_url = _wait_for_doclib(doclib, home / "doclib-home", doclib_log)
            with socket.socket() as listener:
                listener.bind(("127.0.0.1", 0))
                api_port = listener.getsockname()[1]
            origin = f"http://127.0.0.1:{api_port}"
            api_env = os.environ.copy()
            api_env.update(
                {
                    "MINERU_BUSINESS_DB_PATH": str(home / "business.sqlite3"),
                    "MINERU_BUSINESS_UPLOAD_ROOT": str(uploads),
                    "MINERU_BUSINESS_DOCLIB_URL": doclib_url,
                    "MINERU_BUSINESS_MAX_UPLOAD_BYTES": str(5 * 1024 * 1024),
                    "MINERU_BUSINESS_WEB_ROOT": str(WEB_ROOT),
                    "MINERU_BUSINESS_REQUIRE_WEB": "1",
                    "HF_HUB_OFFLINE": "1",
                    "TRANSFORMERS_OFFLINE": "1",
                    "HF_DATASETS_OFFLINE": "1",
                }
            )
            command = (
                "import os, uvicorn; "
                "from mineru.business.api.server import ServerConfig, build_app; "
                "uvicorn.run(build_app(ServerConfig.from_environment()), "
                "host='127.0.0.1', port=int(os.environ['MINERU_TEST_API_PORT']), log_level='warning')"
            )
            api_env["MINERU_TEST_API_PORT"] = str(api_port)
            with api_log.open("wb") as api_stream:
                api = subprocess.Popen(
                    [sys.executable, "-c", command], env=api_env, stdout=api_stream, stderr=subprocess.STDOUT
                )
            try:
                _wait_for_api(api, origin, api_log)
                subprocess.run(
                    [
                        "python3",
                        str(ROOT / "business-web" / "tests" / "live_business_smoke.py"),
                        origin,
                        str(SAMPLE),
                        *(str(path) for path in office_files),
                    ],
                    check=True,
                    timeout=150,
                )
                _verify_live_skill(origin)
            finally:
                _stop(api)
        finally:
            _stop(doclib)
